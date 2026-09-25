import frappe
import openpyxl
import io
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from openpyxl.utils import get_column_letter
from datetime import datetime
from dateutil.relativedelta import relativedelta
from decimal import Decimal, ROUND_HALF_UP

@frappe.whitelist()
def download(from_date=None, to_date=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EL Calculation"

    # Styles
    header_font = Font(bold=True, color="000000")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    title_font = Font(bold=True, size=14)
    subtitle_font = Font(bold=True, size=12)
    thin_side = Side(style="thin", color="000000")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    fill_header = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    fixed_headers = ["S.No", "Emp ID", "Name"]
    sub_headers = ["P", "BH", "NPD"]
    total_sub_headers = ["P", "BH", "NPD"]
    final_summary_headers = [
        "TOTAL",
        "TOTAL EL",
        "EL Carry forward leave",
        "Total EL FOR THE YEAR",
        "EL ENCHASHMENT"
    ]

    dynamic_months = get_months_list(from_date, to_date)
    from_year = datetime.strptime(from_date, "%Y-%m-%d").year

    total_columns = len(fixed_headers) + len(dynamic_months) * len(sub_headers) + len(total_sub_headers) + len(final_summary_headers)
    start_col = 1

    
    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=total_columns)
    cell = ws.cell(row=1, column=start_col, value="TS INTERSEATS INDIA PVT LTD")
    cell.font = title_font
    cell.alignment = center_align
    cell.fill = yellow_fill
    for col in range(start_col, total_columns + 1):
        ws.cell(row=1, column=col).border = thin_border

    ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=total_columns)
    cell = ws.cell(row=2, column=start_col, value=f"EL EARNINGS FOR THE YEAR OF {from_year}")
    cell.font = subtitle_font
    cell.alignment = center_align
    cell.fill = yellow_fill
    for col in range(start_col, total_columns + 1):
        ws.cell(row=2, column=col).border = thin_border

    for idx, header in enumerate(fixed_headers):
        col = start_col + idx
        ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
        for row in range(3, 5):
            cell = ws.cell(row=row, column=col)
            if row == 3: 
                cell.value = header
            cell.font = header_font
            cell.alignment = center_align
            cell.fill = fill_header
            cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = 15

    
    month_start_col = start_col + len(fixed_headers)
    for month in dynamic_months:
        merge_start = month_start_col
        merge_end = month_start_col + len(sub_headers) - 1

        ws.merge_cells(start_row=3, start_column=merge_start, end_row=3, end_column=merge_end)
        month_cell = ws.cell(row=3, column=merge_start, value=month)
        month_cell.font = header_font
        month_cell.alignment = center_align
        month_cell.fill = fill_header
        for col in range(merge_start, merge_end + 1):
            ws.cell(row=3, column=col).border = thin_border
            ws.column_dimensions[get_column_letter(col)].width = 10

        for i, sub in enumerate(sub_headers):
            col = month_start_col + i
            sub_cell = ws.cell(row=4, column=col, value=sub)
            sub_cell.font = header_font
            sub_cell.alignment = center_align
            sub_cell.fill = fill_header
            sub_cell.border = thin_border

        month_start_col += len(sub_headers)

  
    total_col_start = month_start_col
    total_col_end = total_col_start + len(total_sub_headers) - 1

    ws.merge_cells(start_row=3, start_column=total_col_start, end_row=3, end_column=total_col_end)
    total_header_cell = ws.cell(row=3, column=total_col_start, value=f"TOTAL WORKING FOR THE YEAR {from_year}")
    total_header_cell.font = header_font
    total_header_cell.alignment = center_align
    total_header_cell.fill = fill_header
    for col in range(total_col_start, total_col_end + 1):
        ws.cell(row=3, column=col).border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = 10

    for i, sub in enumerate(total_sub_headers):
        col = total_col_start + i
        sub_cell = ws.cell(row=4, column=col, value=sub)
        sub_cell.font = header_font
        sub_cell.alignment = center_align
        sub_cell.fill = fill_header
        sub_cell.border = thin_border

    final_col_start = total_col_end + 1
    for i, header in enumerate(final_summary_headers):
        col = final_col_start + i
        ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
        for row in range(3, 5):
            cell = ws.cell(row=row, column=col)
            if row == 3: 
                cell.value = header
            cell.font = header_font
            cell.alignment = center_align
            cell.fill = fill_header
            cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = 20

    # employees = get_employees(frappe._dict({"from_date": from_date, "to_date": to_date, "employee": None}))
    employees = get_employees(frappe._dict({
    "from_date": from_date,
    "to_date": to_date,
    "employee": frappe.form_dict.get("employee"),
    "employee_catagory": frappe.form_dict.get("employee_catagory"),
    "branch": frappe.form_dict.get("branch"),
    }))
    data_start_row = 5
    for idx, emp in enumerate(employees, start=1):
        row = data_start_row + idx - 1
        base_col = start_col

        ws.cell(row=row, column=base_col, value=idx).alignment = center_align
        ws.cell(row=row, column=base_col).border = thin_border

        ws.cell(row=row, column=base_col + 1, value=emp.name).alignment = center_align
        ws.cell(row=row, column=base_col + 1).border = thin_border

        ws.cell(row=row, column=base_col + 2, value=emp.employee_name).alignment = center_align
        ws.cell(row=row, column=base_col + 2).border = thin_border

        present_map = get_present_counts_by_month(emp.name, from_date, to_date)
        bh_map, dh_map = get_holiday(emp.name, from_date, to_date)

        data_col = base_col + len(fixed_headers)
        for month in dynamic_months:
            for sub in sub_headers:
                val = 0
                if sub == "P":
                    val = present_map.get(month, 0)
                elif sub == "BH":
                    val = bh_map.get(month, 0)
                elif sub == "NPD":
                    val = dh_map.get(month, 0)
                ws.cell(row=row, column=data_col, value=val).alignment = center_align
                ws.cell(row=row, column=data_col).border = thin_border
                data_col += 1


        total_p = sum(present_map.get(month, 0) for month in dynamic_months)
        total_bh = sum(bh_map.get(month, 0) for month in dynamic_months)
        total_npd = sum(dh_map.get(month, 0) for month in dynamic_months)

        for sub in total_sub_headers:
            if sub == "P":
                val = total_p
            elif sub == "BH":
                val = total_bh
            elif sub == "NPD":
                val = total_npd
            else:
                val = 0
        
            ws.cell(row=row, column=data_col, value=val).alignment = center_align
            ws.cell(row=row, column=data_col).border = thin_border
            data_col += 1

        grand_total = total_p + total_bh + total_npd
        # total_el = round(grand_total / 20,0)
        total_el = int(Decimal(grand_total / 20).quantize(0, ROUND_HALF_UP))
        el_carry_forward = frappe.db.get_value('Leave Allocation',{'employee': emp.name,'docstatus': 1,'to_date': ['>=', to_date]},'unused_leaves') or 0

        total_el_for_year = total_el + el_carry_forward
        el_enchashment = max(total_el_for_year - 30, 0)

        summary_values = [
            grand_total,
            total_el,
            el_carry_forward,
            total_el_for_year,
            el_enchashment
        ]

        for val in summary_values:
            ws.cell(row=row, column=data_col, value=val).alignment = center_align
            ws.cell(row=row, column=data_col).border = thin_border
            data_col += 1



    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    frappe.response['filename'] = 'EL Calculation.xlsx'
    frappe.response['filecontent'] = output.getvalue()
    frappe.response['type'] = 'binary'
    frappe.response['doctype'] = None

def get_months_list(from_date_str, to_date_str):
    months = []
    start = datetime.strptime(from_date_str, "%Y-%m-%d")
    end = datetime.strptime(to_date_str, "%Y-%m-%d")
    current = start
    while current <= end:
        months.append(current.strftime("%b-%y"))
        current += relativedelta(months=1)

    return months
    

# def get_employees(filters):
#     conditions = ''
#     if filters.get('employee'):
#         conditions += "AND employee = '%s' " % (filters['employee'])

#     employees = frappe.db.sql("""
#         SELECT name, employee_name, department, date_of_joining
#         FROM tabEmployee
#         WHERE status = 'Active' AND date_of_joining <= %s %s
#     """ % ("%s", conditions), (filters['to_date'],), as_dict=True)

#     left_employees = frappe.db.sql("""
#         SELECT name, employee_name, department, date_of_joining, relieving_date
#         FROM tabEmployee
#         WHERE status = 'Left' AND relieving_date >= %s %s
#     """ % ("%s", conditions), (filters['from_date'],), as_dict=True)

#     employees.extend(left_employees)
#     return employees


def get_employees(filters):
    conditions = []

    if filters.get('employee'):
        conditions.append("name = %(employee)s")

    if filters.get('employee_catagory'):
        conditions.append("employee_catagory = %(employee_catagory)s")

    if filters.get('branch'):
        conditions.append("branch = %(branch)s")

    condition_sql = " AND ".join(conditions)
    if condition_sql:
        condition_sql = " AND " + condition_sql

    employees = frappe.db.sql(f"""
        SELECT name, employee_name, department, date_of_joining
        FROM tabEmployee
        WHERE status = 'Active' AND employment_type != 'Contract' AND date_of_joining <= %(to_date)s {condition_sql}
    """, filters, as_dict=True)

    left_employees = frappe.db.sql(f"""
        SELECT name, employee_name, department, date_of_joining, relieving_date
        FROM tabEmployee
        WHERE status = 'Left' AND employment_type != 'Contract' AND relieving_date >= %(from_date)s {condition_sql}
    """, filters, as_dict=True)

    employees.extend(left_employees)
    return employees


# def get_present_counts_by_month(employee, from_date, to_date):
#     present_records = frappe.db.sql("""
#         SELECT 
#             MONTH(attendance_date) AS month,
#             YEAR(attendance_date) AS year,
#             COUNT(*) AS present_count
#         FROM tabAttendance
#         WHERE
#             docstatus != 2
#             AND employee = %s
#             AND status = 'Present'
#             AND attendance_date BETWEEN %s AND %s
#             AND attendance_date NOT IN (
#                 SELECT holiday_date
#                 FROM `tabHoliday`
#                 WHERE parent = (
#                     SELECT holiday_list FROM `tabEmployee` WHERE name = %s
#                 )
#             )
#         GROUP BY YEAR(attendance_date), MONTH(attendance_date)
#     """, (employee, from_date, to_date, employee), as_dict=True)

#     present_map = {}
#     for row in present_records:
#         label = datetime(row['year'], row['month'], 1).strftime("%b-%y")
#         present_map[label] = row['present_count']

#     return present_map

def get_present_counts_by_month(employee, from_date, to_date):
    present_records = frappe.db.sql("""
                SELECT 
                    MONTH(attendance_date) AS month,
                    YEAR(attendance_date) AS year,
                    SUM(
                        CASE 
                            WHEN status = 'Present' THEN 1
                            WHEN status = 'Half Day' AND shift_status = 'P-OD' THEN 1
                            WHEN status = 'Half Day' THEN 0.5
                            ELSE 0
                        END
                    ) AS present_count
                FROM `tabAttendance`
                WHERE
                    docstatus != 2
                    AND employee = %s
                    AND status IN ('Present', 'Half Day')
                    AND attendance_date BETWEEN %s AND %s
                    AND attendance_date NOT IN (
                        SELECT holiday_date
                        FROM `tabHoliday`
                        WHERE parent = (
                            SELECT holiday_list FROM `tabEmployee` WHERE name = %s
                        )
                    )
                GROUP BY YEAR(attendance_date), MONTH(attendance_date)
            """, (employee, from_date, to_date, employee), as_dict=True)

    present_map = {}
    for row in present_records:
        label = datetime(row['year'], row['month'], 1).strftime("%b-%y")
        present_map[label] = row['present_count']

    return present_map

# def get_holiday(employee, from_date, to_date):
#     holiday_list = frappe.db.get_value("Employee", employee, "holiday_list")
#     if not holiday_list:
#         return {}, {}

#     holiday_records = frappe.db.sql("""
#         SELECT 
#             MONTH(holiday_date) AS month,
#             YEAR(holiday_date) AS year,
#             others, description,
#             COUNT(*) AS count
#         FROM `tabHoliday`
#         WHERE parent = %s
#         AND holiday_date BETWEEN %s AND %s
#         AND TRIM(description) IN ('BLOCK HOLIDAY', 'NON-PRODUCTION DAY')
#         GROUP BY YEAR(holiday_date), MONTH(holiday_date), others
#     """, (holiday_list, from_date, to_date), as_dict=True)

#     bh_map = {}
#     dh_map = {}
#     for row in holiday_records:
#         label = datetime(row['year'], row['month'], 1).strftime("%b-%y")
#         if row['description'] == 'BLOCK HOLIDAY':
#             bh_map[label] = row['count']
#             frappe.log_error(str(row['count']), "BH Count Debug")
#         elif row['description'] == 'NON-PRODUCTION DAY':
#             dh_map[label] = row['count']
    
#     return bh_map, dh_map


def get_holiday(employee, from_date, to_date):
    holiday_list = frappe.db.get_value("Employee", employee, "holiday_list")
    if not holiday_list:
        return {}, {}

    holiday_records = frappe.db.sql("""
        SELECT 
            MONTH(holiday_date) AS month,
            YEAR(holiday_date) AS year,
            SUM(CASE WHEN description = 'BLOCK HOLIDAY' THEN 1 ELSE 0 END) AS bh_count,
            SUM(CASE WHEN description = 'NON-PRODUCTION DAY' THEN 1 ELSE 0 END) AS npd_count
        FROM `tabHoliday`
        WHERE parent = %s
        AND holiday_date BETWEEN %s AND %s
        GROUP BY YEAR(holiday_date), MONTH(holiday_date)
    """, (holiday_list, from_date, to_date), as_dict=True)

    bh_map = {}
    dh_map = {}

    for row in holiday_records:
        label = datetime(row['year'], row['month'], 1).strftime("%b-%y")
        bh_map[label] = row['bh_count'] or 0
        dh_map[label] = row['npd_count'] or 0

    return bh_map, dh_map
