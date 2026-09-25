import frappe
import openpyxl
import io
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from openpyxl.utils import get_column_letter
from datetime import datetime
from dateutil.relativedelta import relativedelta

@frappe.whitelist()
def download(from_date=None, to_date=None, employee=None, employee_catagory=None, branch=None):
    if not from_date or not to_date:
        frappe.throw("From Date and To Date are mandatory")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Yearly Bonus Calculation"


    header_font = Font(bold=True, color="000000")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    title_font = Font(bold=True, size=14)
    subtitle_font = Font(bold=True, size=12)
    thin_side = Side(style="thin", color="000000")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    fill_header = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    
    
    dynamic_months = get_months_list(from_date, to_date)
    from_year = datetime.strptime(from_date, "%Y-%m-%d").year
    to_year = datetime.strptime(to_date, "%Y-%m-%d").year
    
    fixed_headers = ["S.No", "Emp ID", "Emp Name"]
    final_summary_headers = [
        "TOTAL",
        f"Paid Days {from_year} to {to_year}",
        "Position Bonus",
        f"Bonus Payable Salary {from_year} to {to_year}",
        f"Actual Payable Bonus {to_year}",
        "Remarks"
    ]
    total_columns = len(fixed_headers) + len(dynamic_months) + len(final_summary_headers)
    start_col = 1

 
    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=total_columns)
    cell = ws.cell(row=1, column=start_col, value="TS INTERSEATS INDIA PVT LTD")
    cell.font = title_font
    cell.alignment = center_align
    cell.fill = yellow_fill
    for col in range(start_col, total_columns + 1):
        ws.cell(row=1, column=col).border = thin_border

  
    ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=total_columns)
    cell = ws.cell(row=2, column=start_col, value=f"STAFF PAYMENT OF BONUS FROM {from_year} to {to_year}")
    cell.font = subtitle_font
    cell.alignment = center_align
    cell.fill = yellow_fill
    for col in range(start_col, total_columns + 1):
        ws.cell(row=2, column=col).border = thin_border

   
    for idx, header in enumerate(fixed_headers):
        col = start_col + idx
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.alignment = center_align
        cell.fill = fill_header
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = 15

    
    month_start_col = start_col + len(fixed_headers)

    for month in dynamic_months:
        month_cell = ws.cell(row=3, column=month_start_col, value=month)
        month_cell.font = header_font
        month_cell.alignment = center_align
        month_cell.fill = fill_header
        month_cell.border = thin_border  
        ws.column_dimensions[get_column_letter(month_start_col)].width = 12

        month_start_col += 1

    for i, header in enumerate(final_summary_headers):
        col = month_start_col + i
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.alignment = center_align
        cell.fill = fill_header
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = 20

    # employees = get_employees(frappe._dict({"from_date": from_date, "to_date": to_date, "employee": None}))
    employees = get_employees(frappe._dict({
    "from_date": from_date,
    "to_date": to_date,
    "employee": frappe.form_dict.get("employee"),
    "employee_catagory": frappe.form_dict.get("employee_catagory"),
    "branch": frappe.form_dict.get("branch"),
    }))

    data_start_row = 4
    for idx, emp in enumerate(employees, start=1):
        row = data_start_row + idx - 1
        base_col = start_col

        cell = ws.cell(row=row, column=base_col, value=idx)
        cell.alignment = center_align
        cell.border = thin_border

        cell = ws.cell(row=row, column=base_col + 1, value=emp.name)
        cell.alignment = center_align
        cell.border = thin_border

        cell = ws.cell(row=row, column=base_col + 2, value=emp.employee_name)
        cell.alignment = center_align
        cell.border = thin_border

        pay_days_count = get_paid_days(emp.name, from_date, to_date)
        data_col = base_col + len(fixed_headers)

        row_total =0

        for month in dynamic_months:
            val = pay_days_count.get(month, 0)
            row_total += val
            cell = ws.cell(row=row, column=data_col, value=val)
            cell.alignment = center_align
            cell.border = thin_border
            data_col += 1
        
        cell = ws.cell(row=row, column=data_col, value=row_total)
        cell.alignment = center_align
        cell.border = thin_border
        data_col += 1

        cell = ws.cell(row=row, column=data_col, value=row_total)
        cell.alignment = center_align
        cell.border = thin_border
        data_col += 1

        emp_data = get_employee_data(emp.name)

        cell = ws.cell(row=row, column=data_col, value=emp_data["position_bonus"])
        cell.alignment = center_align
        cell.border = thin_border
        data_col += 1

        cell = ws.cell(row=row, column=data_col, value=emp_data["bonus_payable"])
        cell.alignment = center_align
        cell.border = thin_border
        data_col += 1

        cell = ws.cell(row=row, column=data_col, value=round((emp_data["bonus_payable"]/365)*row_total))
        cell.alignment = center_align
        cell.border = thin_border
        data_col += 1

        cell = ws.cell(row=row, column=data_col)
        cell.alignment = center_align
        cell.border = thin_border
        data_col += 1

    #Add Total Row at Bottom 
    total_row = data_start_row + len(employees)

    # Merge TOTAL cell across 3 columns (S.No, Emp ID, Emp Name)
    ws.merge_cells(start_row=total_row, start_column=start_col, end_row=total_row, end_column=start_col+2)
    total_cell = ws.cell(row=total_row, column=start_col, value="TOTAL")
    total_cell.font = header_font
    total_cell.alignment = center_align

    # Apply border to merged TOTAL cell (all merged columns)
    for c in range(start_col, start_col+3):
        ws.cell(row=total_row, column=c).border = thin_border

    # Loop through columns after Emp Name up to (but excluding Remarks)
    start_data_row = data_start_row
    end_data_row = total_row - 1
    last_col = ws.max_column 

    for c in range(start_col + 3, last_col):
        col_letter = get_column_letter(c)
        sum_formula = f"=SUM({col_letter}{start_data_row}:{col_letter}{end_data_row})"
        total_cell = ws.cell(row=total_row, column=c, value=sum_formula)
        total_cell.font = header_font
        total_cell.alignment = center_align
        total_cell.border = thin_border

    # Keep Remarks column empty but with border
    remarks_cell = ws.cell(row=total_row, column=last_col, value="")
    remarks_cell.border = thin_border

    
    xlsx_file = io.BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)

   
    frappe.response['filename'] = f"Yearly_Bonus_Calculation.xlsx"
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'


def get_months_list(from_date_str, to_date_str):
    months = []
    start = datetime.strptime(from_date_str, "%Y-%m-%d")
    end = datetime.strptime(to_date_str, "%Y-%m-%d")
    current = start
    while current <= end:
        months.append(current.strftime("%b-%y"))
        current += relativedelta(months=1)
    return months


# def get_employees(filters):
#     conditions = ''
#     if filters.get('employee'):
#         conditions += "AND employee = '%s' " % (filters['employee'])

#     employees = frappe.db.sql("""
#         SELECT name, employee_name, department, date_of_joining
#         FROM tabEmployee
#         WHERE status = 'Active' AND employment_type != 'Contract' AND date_of_joining <= %s %s
#     """ % ("%s", conditions), (filters['to_date'],), as_dict=True)

#     left_employees = frappe.db.sql("""
#         SELECT name, employee_name, department, date_of_joining, relieving_date
#         FROM tabEmployee
#         WHERE status = 'Left' AND  employment_type != 'Contract' AND relieving_date >= %s %s
#     """ % ("%s", conditions), (filters['from_date'],), as_dict=True)

#     employees.extend(left_employees)
#     return employees

def get_employees(filters):
    conditions = []

    if filters.get('employee'):
        conditions.append("name = %(employee)s")

    if filters.get('employee_catagory'):
        conditions.append("employee_catagory = %(employee_catagory)s")

    if filters.get('branch'):
        conditions.append("branch = %(branch)s")

    condition_sql = " AND ".join(conditions)
    if condition_sql:
        condition_sql = " AND " + condition_sql

    employees = frappe.db.sql(f"""
        SELECT name, employee_name, department, date_of_joining
        FROM tabEmployee
        WHERE status = 'Active' AND employment_type != 'Contract' AND date_of_joining <= %(to_date)s {condition_sql}
    """, filters, as_dict=True)

    left_employees = frappe.db.sql(f"""
        SELECT name, employee_name, department, date_of_joining, relieving_date
        FROM tabEmployee
        WHERE status = 'Left' AND employment_type != 'Contract' AND relieving_date >= %(from_date)s {condition_sql}
    """, filters, as_dict=True)

    employees.extend(left_employees)
    return employees

def get_paid_days(employee, from_date, to_date):
    pay_day = frappe.db.sql("""
        SELECT  
            MONTH(end_date) AS month,
            YEAR(end_date) AS year,
            SUM(payment_days) AS payment_days
        FROM `tabSalary Slip`
        WHERE docstatus = 1
          AND employee = %s
          AND end_date BETWEEN %s AND %s
        GROUP BY YEAR(end_date), MONTH(end_date)
    """, (employee, from_date, to_date), as_dict=True)

    pay_days_count = {}
    for row in pay_day:
        
        label = datetime(row['year'], row['month'], 1).strftime("%b-%y")
        pay_days_count[label] = row['payment_days']


    return pay_days_count


def get_employee_data(employee):
    data = frappe.db.get_value(
        "Employee",
        {"employee": employee},
        ["gross", "position_bonus"],
        as_dict=True
    )
    if not data:
        return {"gross": 0, "position_bonus": 0}

    return {
        "position_bonus": data.position_bonus or 0,
        "bonus_payable": (data.gross or 0) - (data.position_bonus or 0)
    }

