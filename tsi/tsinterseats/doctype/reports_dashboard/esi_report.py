from __future__ import unicode_literals
import frappe
from frappe.utils import cstr, add_days, date_diff, getdate, format_date
from frappe import _, bold
from frappe.utils.csvutils import UnicodeWriter, read_csv_content
from frappe.utils.data import format_date
from frappe.utils.file_manager import get_file
from frappe.model.document import Document
from frappe.utils.background_jobs import enqueue

from datetime import date, timedelta, datetime
import openpyxl
from openpyxl import Workbook
import re
from frappe import _
import frappe
from frappe.model.document import Document
from datetime import date, timedelta, datetime,time
from frappe.utils import (getdate, cint, add_months, date_diff, add_days,
	nowdate, get_datetime_str, cstr, get_datetime, now_datetime, format_datetime,today, format_date)
import math
from frappe.utils import add_months, cint, flt, getdate, time_diff_in_hours,time_diff_in_seconds
import locale


import openpyxl
import xlrd
import re
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
from six import BytesIO, string_types
import openpyxl.styles as styles

@frappe.whitelist()
def download():
	filename = 'ESI Report'
	test = build_xlsx_response(filename)
	
def make_xlsx(data, sheet_name=None, wb=None, column_widths=None):
	args = frappe.local.form_dict
	column_widths = column_widths or []
	if wb is None:
		wb = openpyxl.Workbook()
		 
	ws = wb.create_sheet(sheet_name, 0)
	ws.column_dimensions['A'].width = 15
	ws.column_dimensions['B'].width = 18 
	ws.column_dimensions['C'].width = 18 
	ws.column_dimensions['D'].width = 20
	ws.column_dimensions['E'].width = 28
	ws.column_dimensions['F'].width = 18

	thin_border = Border(
	left=Side(style='thin'),
	right=Side(style='thin'),
	top=Side(style='thin'),
	bottom=Side(style='thin')
	)
	start_date = args.start_date
	month_year = datetime.strptime(start_date, "%Y-%m-%d").strftime("%B-%Y").upper()

	ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
	ws.cell(row=1, column=1).value = f"ESI CONTRIBUTION FOR THE MONTH {month_year}"
	ws.cell(row=1, column=1).font = Font(bold=True)
	ws.cell(row=1, column=1).alignment = Alignment(horizontal='center', vertical='center')

	ws.cell(row=2, column=1).value = "ESI contribution of Workers"
	ws.cell(row=2, column=1).font = Font(bold=True)
	ws.cell(row=2, column=1).alignment = Alignment(horizontal='left', vertical='center')

	ws.append(["Employee ID","Employee Name","Gross Salary","EMPE'S ESI","EMPLR'S ESI"])
	align_center = Alignment(horizontal='center',vertical='center')
	for header in ws.iter_rows(min_row=1, max_row=3, min_col=1, max_col=5):
		for cell in header:
			cell.font = Font(bold=True)
			cell.alignment = align_center
			cell.border = thin_border

	data1= get_data(args)
	for row in data1:
		ws.append(row)

	total_gross = 0
	total_esi = 0
	total_employer_esi = 0

	for r in data1:
		total_gross += float(r[2] or 0)
		total_esi += float(r[3] or 0)
		total_employer_esi += float(r[4] or 0)

	ws.append([
		"TOTAL",
		"",
		round(total_gross, 2),
		round(total_esi, 2),
		round(total_employer_esi, 2)
	])

	total_row = ws.max_row

	for col in range(1, 6):
		cell = ws.cell(row=total_row, column=col)
		cell.font = Font(bold=True)
		cell.alignment = Alignment(horizontal='center', vertical='center')
		cell.border = thin_border

	max_row = ws.max_row
	max_col = ws.max_column
	for row in ws.iter_rows(min_row=2, max_row=max_row, min_col=1, max_col=max_col):
		for cell in row:
			cell.border = thin_border
			cell.alignment = Alignment(vertical='center')

	xlsx_file = BytesIO()
	wb.save(xlsx_file)
	return xlsx_file

def build_xlsx_response(filename):
	xlsx_file = make_xlsx(filename)
	frappe.response['filename'] = filename + '.xlsx'
	frappe.response['filecontent'] = xlsx_file.getvalue()
	frappe.response['type'] = 'binary'

@frappe.whitelist()
def get_data(args):
	data = []
	row = []
	# salary_slips = frappe.db.sql("""select * from `tabSalary Slip` where start_date between '%s' and '%s' and docstatus!=2"""%(args.start_date, args.end_date), as_dict=True)
	salary_slips = frappe.db.sql("""
		SELECT
			ss.employee,
			ss.employee_name,
			ss.gross_pay,
			e.is_esic_eligible
		FROM `tabSalary Slip` ss
		INNER JOIN `tabEmployee` e ON e.name = ss.employee
		WHERE ss.start_date BETWEEN %s AND %s
		AND ss.docstatus != 2
		AND ss.gross_pay > 0
		AND (ss.gross_pay < 21000 OR e.is_esic_eligible = 1)
	""", (args.start_date, args.end_date), as_dict=True)

	for ss in salary_slips:
		emp_id = ss.employee
		emp_name = ss.employee_name
		gs = float(ss.get("gross_pay") or 0)
		esi_amount= round(gs*0.0075,0)
		eplyr_esi = round(gs * 0.0325, 0) if esi_amount > 0 else 0
		# esi_amount = frappe.get_value('Salary Detail', {'salary_component': "ESI",'parent': ss.name,'docstatus': ['!=', 2]}, ["amount"] or 0)
		row = [emp_id or "-",emp_name or "-",gs or "-",esi_amount or 0,eplyr_esi or 0]
		data.append(row)
	return data